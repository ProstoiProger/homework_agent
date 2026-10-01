import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from db import init_database
from rag_service import ResourceService
from services import DomainError


class ResourceServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.database = init_database(f"sqlite:///{(root / 'test.db').as_posix()}")
        self.resources = ResourceService(self.database, root / "resources")

    def tearDown(self):
        self.database.engine.dispose()
        self.temp_dir.cleanup()

    def test_ingest_search_deduplicate_and_delete_text_resource(self):
        data = (
            "Производная функции показывает скорость её изменения.\n"
            "Для степенной функции используется правило дифференцирования.\n"
            "Пример: производная x в квадрате равна двум x."
        ).encode("utf-8")

        created = self.resources.ingest_bytes(data, "Алгебра.txt", "text/plain", "tg-1")
        duplicate = self.resources.ingest_bytes(data, "Копия.txt", "text/plain", "tg-2")
        results = self.resources.search("задачи по производным", limit=3)

        self.assertFalse(created["already_exists"])
        self.assertTrue(duplicate["already_exists"])
        self.assertEqual(len(self.resources.list_documents()), 1)
        self.assertEqual(results[0]["filename"], "Алгебра.txt")
        self.assertIn("Производная", results[0]["text"])

        deleted = self.resources.delete_document(created["id"])
        self.assertEqual(deleted["filename"], "Алгебра.txt")
        self.assertEqual(self.resources.list_documents(), [])

    def test_ingests_docx(self):
        document = Document()
        document.add_paragraph("Квадратное уравнение решается через дискриминант.")
        document.add_paragraph("В учебнике приведены упражнения и примеры решения.")
        buffer = BytesIO()
        document.save(buffer)

        created = self.resources.ingest_bytes(
            buffer.getvalue(), "Учебник.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        self.assertGreater(created["char_count"], 20)
        self.assertEqual(self.resources.search("дискриминант")[0]["filename"], "Учебник.docx")

    def test_rejects_unsupported_file(self):
        with self.assertRaisesRegex(DomainError, "PDF, DOCX, TXT"):
            self.resources.ingest_bytes(b"some data", "archive.zip")


if __name__ == "__main__":
    unittest.main()
