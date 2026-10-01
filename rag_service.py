from __future__ import annotations

import hashlib
import math
import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader
from sqlalchemy import select

from db import Database
from models import ResourceChunk, ResourceDocument
from services import DomainError


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
WORD_RE = re.compile(r"[0-9A-Za-zА-Яа-яЁё]{2,}", re.UNICODE)


def _tokens(text: str) -> list[str]:
    result = []
    for token in WORD_RE.findall(text):
        normalized = token.casefold().replace("ё", "е")
        result.append(normalized)
        if len(normalized) >= 7:
            result.append(f"^{normalized[:6]}")
    return result


def _safe_filename(filename: str) -> str:
    name = Path(filename).name.strip() or "resource"
    stem = re.sub(r"[^0-9A-Za-zА-Яа-яЁё._ -]+", "_", Path(name).stem).strip(" ._")
    return f"{(stem or 'resource')[:120]}{Path(name).suffix.lower()}"


def _chunk_sections(
    sections: list[tuple[str, str]], max_chars: int = 1800, overlap: int = 200
) -> list[tuple[str, str]]:
    chunks: list[tuple[str, str]] = []
    for location, raw_text in sections:
        text = re.sub(r"[ \t]+", " ", raw_text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        start = 0
        while start < len(text):
            end = min(start + max_chars, len(text))
            if end < len(text):
                boundary = max(
                    text.rfind("\n", start + max_chars // 2, end),
                    text.rfind(". ", start + max_chars // 2, end),
                    text.rfind(" ", start + max_chars // 2, end),
                )
                if boundary > start:
                    end = boundary + 1
            chunk = text[start:end].strip()
            if chunk:
                chunks.append((location, chunk))
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks


class ResourceService:
    def __init__(self, database: Database, storage_dir: Path):
        self.database = database
        self.storage_dir = storage_dir.resolve()
        self.storage_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _extract(data: bytes, extension: str) -> list[tuple[str, str]]:
        try:
            if extension == ".pdf":
                reader = PdfReader(BytesIO(data))
                if reader.is_encrypted:
                    raise DomainError("PDF защищён паролем — такой файл прочитать нельзя.")
                return [
                    (f"стр. {number}", page.extract_text() or "")
                    for number, page in enumerate(reader.pages, start=1)
                ]
            if extension == ".docx":
                document = Document(BytesIO(data))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                return [("документ", text)]
            if extension in {".txt", ".md"}:
                try:
                    text = data.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = data.decode("cp1251")
                return [("текст", text)]
        except DomainError:
            raise
        except Exception as exc:
            raise DomainError(f"Не удалось прочитать файл: {exc}") from exc
        raise DomainError("Поддерживаются только PDF, DOCX, TXT и MD.")

    def ingest_bytes(
        self,
        data: bytes,
        filename: str,
        mime_type: str | None = None,
        telegram_file_id: str | None = None,
    ) -> dict[str, Any]:
        safe_name = _safe_filename(filename)
        extension = Path(safe_name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise DomainError("Поддерживаются только PDF, DOCX, TXT и MD.")
        if not data:
            raise DomainError("Файл пустой.")

        digest = hashlib.sha256(data).hexdigest()
        with self.database.session() as session:
            existing = session.scalar(
                select(ResourceDocument).where(ResourceDocument.sha256 == digest)
            )
            if existing:
                return {**self._document_dict(existing), "already_exists": True}

        sections = self._extract(data, extension)
        chunks = _chunk_sections(sections)
        char_count = sum(len(text) for _, text in sections)
        if not chunks or char_count < 20:
            raise DomainError(
                "В файле не найден читаемый текст. Если это скан, сначала нужен OCR."
            )

        stored_name = f"{digest[:16]}_{safe_name}"
        stored_path = (self.storage_dir / stored_name).resolve()
        if not stored_path.is_relative_to(self.storage_dir):
            raise DomainError("Недопустимое имя файла.")
        stored_path.write_bytes(data)

        try:
            with self.database.transaction() as session:
                document = ResourceDocument(
                    filename=safe_name,
                    mime_type=mime_type,
                    telegram_file_id=telegram_file_id,
                    sha256=digest,
                    storage_path=str(stored_path),
                    char_count=char_count,
                    chunk_count=len(chunks),
                )
                document.chunks = [
                    ResourceChunk(chunk_index=index, location=location, text=text)
                    for index, (location, text) in enumerate(chunks, start=1)
                ]
                session.add(document)
                session.flush()
                return {**self._document_dict(document), "already_exists": False}
        except Exception:
            stored_path.unlink(missing_ok=True)
            raise

    @staticmethod
    def _document_dict(document: ResourceDocument) -> dict[str, Any]:
        return {
            "id": document.id,
            "filename": document.filename,
            "char_count": document.char_count,
            "chunk_count": document.chunk_count,
            "created_at": document.created_at.isoformat() if document.created_at else None,
        }

    def list_documents(self) -> list[dict[str, Any]]:
        with self.database.session() as session:
            documents = session.scalars(
                select(ResourceDocument).order_by(
                    ResourceDocument.created_at.desc(), ResourceDocument.id.desc()
                )
            ).all()
            return [self._document_dict(document) for document in documents]

    def delete_document(self, document_id: int) -> dict[str, Any]:
        with self.database.transaction() as session:
            document = session.get(ResourceDocument, document_id)
            if not document:
                raise DomainError("Ресурс не найден.")
            payload = self._document_dict(document)
            stored_path = Path(document.storage_path).resolve()
            session.delete(document)
        if stored_path.is_relative_to(self.storage_dir):
            stored_path.unlink(missing_ok=True)
        return payload

    def search(
        self, query: str, resource_id: int | None = None, limit: int = 5
    ) -> list[dict[str, Any]]:
        query_terms = _tokens(query)
        if not query_terms:
            raise DomainError("Нужна содержательная тема для поиска по ресурсам.")
        limit = max(1, min(int(limit), 8))

        with self.database.session() as session:
            statement = (
                select(ResourceChunk, ResourceDocument)
                .join(ResourceDocument, ResourceDocument.id == ResourceChunk.document_id)
                .order_by(ResourceDocument.id, ResourceChunk.chunk_index)
            )
            if resource_id is not None:
                statement = statement.where(ResourceDocument.id == resource_id)
            rows = session.execute(statement).all()
        if not rows:
            raise DomainError("Загруженных ресурсов для поиска нет.")

        tokenized = [_tokens(chunk.text) for chunk, _ in rows]
        document_count = len(tokenized)
        average_length = sum(len(tokens) for tokens in tokenized) / document_count
        document_frequency = Counter()
        for tokens in tokenized:
            document_frequency.update(set(tokens))

        scores: list[tuple[float, int]] = []
        for index, tokens in enumerate(tokenized):
            frequencies = Counter(tokens)
            length = len(tokens)
            score = 0.0
            for term in set(query_terms):
                frequency = frequencies.get(term, 0)
                if not frequency:
                    continue
                frequency_in_docs = document_frequency[term]
                inverse_frequency = math.log(
                    1 + (document_count - frequency_in_docs + 0.5) / (frequency_in_docs + 0.5)
                )
                denominator = frequency + 1.5 * (
                    1 - 0.75 + 0.75 * length / max(average_length, 1)
                )
                score += inverse_frequency * frequency * 2.5 / denominator
            if score > 0:
                scores.append((score, index))

        results = []
        for score, row_index in sorted(scores, reverse=True)[:limit]:
            chunk, document = rows[row_index]
            results.append(
                {
                    "resource_id": document.id,
                    "filename": document.filename,
                    "location": chunk.location,
                    "chunk": chunk.chunk_index,
                    "score": round(score, 4),
                    "text": chunk.text,
                }
            )
        return results
